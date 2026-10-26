"""Экспорт статьи в формат docx (НФТ-06): HTML статьи переводится в абзацы Word."""
import base64
import binascii
import io
from html.parser import HTMLParser

from docx import Document
from docx.shared import Pt

BLOCKS = {'p', 'div', 'li', 'blockquote', 'pre', 'h1', 'h2', 'h3', 'h4', 'td', 'th'}
HEADINGS = {'h1': 1, 'h2': 2, 'h3': 3, 'h4': 4}


class _Builder(HTMLParser):
    def __init__(self, document):
        super().__init__(convert_charrefs=True)
        self.doc = document
        self.paragraph = None
        self.style = None
        self.bold = self.italic = self.underline = self.code = 0
        self.lists = []  # стек открытых списков: {'tag': 'ul'|'ol', 'num_id': ...}
        self.li_paragraph = None  # абзац пункта списка, ещё без текста
        self.table = None
        self.cell = None
        self.row = None

    def _new_paragraph(self, style=None):
        target = self.cell if self.cell is not None else self.doc
        self.paragraph = target.add_paragraph(style=style)
        return self.paragraph

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if self.cell is not None and (tag in BLOCKS or tag in HEADINGS or tag in ('br', 'hr')):
            # внутри ячейки таблицы абзацев нет: блоки превращаются в переносы строк
            if self.cell and self.cell[-1] != '\n':
                self.cell.append('\n')
            if tag == 'li':
                self.cell.append('• ')
            return
        if tag in HEADINGS:
            self.paragraph = self.doc.add_heading(level=HEADINGS[tag])
        elif tag in ('ul', 'ol'):
            self.lists.append({'tag': tag, 'num_id': None})
            self.paragraph = self.li_paragraph = None
        elif tag == 'li':
            self._list_item()
        elif tag == 'blockquote':
            self._new_paragraph('Quote')
        elif tag == 'pre':
            self._new_paragraph()
            self.code += 1
        elif tag in ('p', 'div'):
            # <li><p>текст</p></li>: текст идёт в абзац пункта, а не в новый абзац под маркером
            if self.li_paragraph is not None and self.li_paragraph is self.paragraph:
                self.li_paragraph = None
            elif self.lists:
                self._new_paragraph('List Continue' if len(self.lists) == 1 else 'List Continue 2')
            else:
                self._new_paragraph()
        elif tag == 'br':
            if self.paragraph is None:
                self._new_paragraph()
            self.paragraph.add_run().add_break()
        elif tag == 'hr':
            self._new_paragraph().add_run('—' * 20)
        elif tag in ('strong', 'b'):
            self.bold += 1
        elif tag in ('em', 'i'):
            self.italic += 1
        elif tag == 'u':
            self.underline += 1
        elif tag == 'code':
            self.code += 1
        elif tag == 'table':
            self.table = {'rows': []}
        elif tag == 'tr' and self.table is not None:
            self.row = []
            self.table['rows'].append(self.row)
        elif tag in ('td', 'th') and self.row is not None:
            self.cell = ['']
            self.row.append(self.cell)
        elif tag == 'img':
            self._image(attrs.get('src', ''))
        elif tag == 'a' and attrs.get('href'):
            self._link = attrs['href']

    def handle_endtag(self, tag):
        if tag in ('ul', 'ol') and self.lists:
            self.lists.pop()
            self.paragraph = self.li_paragraph = None
        elif tag in ('strong', 'b'):
            self.bold = max(0, self.bold - 1)
        elif tag in ('em', 'i'):
            self.italic = max(0, self.italic - 1)
        elif tag == 'u':
            self.underline = max(0, self.underline - 1)
        elif tag in ('code', 'pre'):
            self.code = max(0, self.code - 1)
        elif tag in ('td', 'th'):
            self.cell = None
        elif tag == 'table':
            self._flush_table()
        if tag in BLOCKS or tag in HEADINGS:
            self.paragraph = None

    def _list_item(self):
        current = self.lists[-1] if self.lists else {'tag': 'ul', 'num_id': None}
        depth = min(len(self.lists), 3)
        base = 'List Number' if current['tag'] == 'ol' else 'List Bullet'
        style = base if depth <= 1 else f'{base} {depth}'
        paragraph = self._new_paragraph(style)
        if current['tag'] == 'ol':
            # каждый нумерованный список начинается с 1, а не продолжает предыдущий
            if current['num_id'] is None:
                current['num_id'] = self._restarted_num(paragraph)
            if current['num_id'] is not None:
                num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_ilvl().val = 0
                num_pr.get_or_add_numId().val = current['num_id']
        self.li_paragraph = paragraph

    def _restarted_num(self, paragraph):
        try:
            numbering = self.doc.part.numbering_part.numbering_definitions._numbering
            style_num_id = paragraph.style.element.pPr.numPr.numId.val
            abstract_id = numbering.num_having_numId(style_num_id).abstractNumId.val
            num = numbering.add_num(abstract_id)
            num.add_lvlOverride(ilvl=0).add_startOverride(1)
            return num.numId
        except (AttributeError, KeyError):
            return None

    def _flush_table(self):
        table, self.table, self.row = self.table, None, None
        rows = [r for r in (table or {}).get('rows', []) if r]
        if not rows:
            return
        width = max(len(r) for r in rows)
        out = self.doc.add_table(rows=len(rows), cols=width)
        out.style = 'Table Grid'
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                out.cell(i, j).text = ''.join(cell).strip()

    def _image(self, src):
        if self.cell is not None:
            self.cell.append('[изображение]')
            return
        if not src.startswith('data:image/') or ',' not in src:
            if src:
                self._write(f'[изображение: {src}]')
            return
        try:
            raw = base64.b64decode(src.split(',', 1)[1])
            if self.paragraph is None:
                self._new_paragraph()
            from docx.shared import Cm
            self.paragraph.add_run().add_picture(io.BytesIO(raw), width=Cm(14))
        except (binascii.Error, ValueError, OSError, Exception):
            self._write('[изображение]')

    def _write(self, text):
        if self.cell is not None and self.row is not None:
            self.cell.append(text)
            return
        if self.paragraph is None:
            self._new_paragraph()
        run = self.paragraph.add_run(text)
        run.bold = bool(self.bold) or None
        run.italic = bool(self.italic) or None
        run.underline = bool(self.underline) or None
        if self.code:
            run.font.name = 'Courier New'
            run.font.size = Pt(10)

    def handle_data(self, data):
        if not data.strip() and self.cell is None and not self.code and (
                self.paragraph is None or not self.paragraph.runs):
            return
        self._write(data if self.code else ' '.join(data.split()) if data.strip() else ' ')


def article_to_docx(article):
    """Возвращает байты docx-файла с заголовком и содержимым статьи."""
    document = Document()
    document.add_heading(article.title, level=0)
    meta = [f'Тип: {article.get_article_type_display()}']
    if article.section:
        meta.append(f'Раздел: {article.section.name}')
    if article.scope:
        meta.append(f'Область применения: {article.scope}')
    meta.append(f'Версия {article.version}, обновлено {article.updated_at:%d.%m.%Y}')
    if article.last_verified_at:
        meta.append(f'Актуальность подтверждена {article.last_verified_at:%d.%m.%Y}')
    document.add_paragraph('; '.join(meta))
    builder = _Builder(document)
    builder.feed(article.content or '')
    builder.close()
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
